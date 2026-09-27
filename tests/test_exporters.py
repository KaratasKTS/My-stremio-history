import csv
from datetime import date

import openpyxl
import pytest

from stremio_watched.exporters import FORMAT_ORDER, default_filename, export, parse_formats
from stremio_watched.exporters.columns import HEADERS
from stremio_watched.exporters.terminal import print_table
from stremio_watched.exporters.txt import write_txt
from stremio_watched.exporters.xlsx import write_xlsx
from stremio_watched.watched import attach_meta, extract_watched


@pytest.fixture
def movies(items, metas):
    return attach_meta(extract_watched(items), metas)


def test_xlsx_layout(movies, tmp_path):
    path = tmp_path / "out.xlsx"
    write_xlsx(movies, path)
    ws = openpyxl.load_workbook(path).active
    assert tuple(c.value for c in ws[1]) == HEADERS
    assert ws.max_row == len(movies) + 1
    assert ws["A2"].value == "Finished not in library"  # newest first; Stremio's own title is kept
    assert ws["B2"].value == "1994" and ws["C2"].value == "Drama"  # metadata fills the rest
    assert ws["E2"].number_format == "0%" and 0.9 < ws["E2"].value <= 1
    assert ws["P2"].hyperlink.target == "https://www.imdb.com/title/tt0111161/"
    assert ws.freeze_panes == "A2"


def test_csv_matches_xlsx_columns(movies, tmp_path):
    written = export(movies, ["csv"], tmp_path, "someone@example.com")
    assert written == [tmp_path / default_filename("someone@example.com", "csv")]
    with written[0].open(newline="", encoding="utf-8") as fh:
        rows = list(csv.reader(fh))
    assert tuple(rows[0]) == HEADERS
    assert len(rows) == len(movies) + 1
    assert rows[1][0] == "Finished not in library" and rows[1][4] == "93%"


def test_txt_is_chronological_one_per_line(movies, tmp_path):
    path = tmp_path / "out.txt"
    write_txt(movies, path)
    lines = path.read_text(encoding="utf-8").splitlines()
    assert len(lines) == len(movies)
    assert lines[0].endswith("| No imdb id | In progress 50%")
    assert lines[-1].endswith("| Finished not in library (1994) | Watched 93%")
    dates = [line.split(" | ")[0] for line in lines]
    assert dates == sorted(dates)


def test_terminal_table(movies, capsys):
    print_table(movies)
    out = capsys.readouterr().out
    lines = [line for line in out.splitlines() if line.strip()]
    assert lines[0].split()[:2] == ["#", "Title"]
    assert len(lines) == len(movies) + 2  # header + rule + rows
    assert "Finished not in library" in lines[2] and lines[2].rstrip().endswith("No")


def test_print_table_handles_empty(capsys):
    print_table([])
    assert "Title" in capsys.readouterr().out


def test_export_all_formats(movies, tmp_path, capsys):
    written = export(movies, parse_formats("all"), tmp_path / "sub", "user.name+tag@example.com")
    assert [p.suffix for p in written] == [".xlsx", ".csv", ".txt"]
    assert all(p.name.startswith("stremio_watched_user.name_tag_") for p in written)
    assert "Finished not in library" in capsys.readouterr().out  # terminal was part of "all"


def test_parse_formats():
    assert parse_formats("all") == list(FORMAT_ORDER)
    assert parse_formats("") == list(FORMAT_ORDER)
    assert parse_formats("txt, XLSX") == ["xlsx", "txt"]
    with pytest.raises(ValueError, match="unknown format"):
        parse_formats("pdf")


def test_default_filename():
    assert default_filename("me@example.com", "txt", date(2026, 9, 27)) == "stremio_watched_me_2026-09-27.txt"
    assert default_filename("@weird", "csv", date(2026, 1, 1)) == "stremio_watched_account_2026-01-01.csv"
