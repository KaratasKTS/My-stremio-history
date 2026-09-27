"""Excel workbook via openpyxl, newest first."""

from __future__ import annotations

from collections.abc import Sequence
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..models import WatchedMovie
from ..watched import sort_newest_first
from .columns import COLUMNS, column_index, row_values

_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor="4F4F4F")
_LINK_FONT = Font(color="0563C1", underline="single")


def write_xlsx(movies: Sequence[WatchedMovie], path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Watched Movies"

    for col, column in enumerate(COLUMNS, 1):
        cell = ws.cell(row=1, column=col, value=column.header)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(col)].width = column.width

    progress_col = column_index("Watched %")
    date_col = column_index("Last watched")
    url_col = column_index("IMDb URL")
    ordered = sort_newest_first(movies)
    for row, movie in enumerate(ordered, 2):
        for col, value in enumerate(row_values(movie), 1):
            ws.cell(row=row, column=col, value=value)
        ws.cell(row=row, column=progress_col).number_format = "0%"
        ws.cell(row=row, column=date_col).number_format = "yyyy-mm-dd hh:mm"
        if movie.imdb_url:
            link = ws.cell(row=row, column=url_col)
            link.hyperlink = movie.imdb_url
            link.font = _LINK_FONT

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{len(ordered) + 1}"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
